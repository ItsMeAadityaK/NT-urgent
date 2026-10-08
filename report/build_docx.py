"""
build_docx.py: the same report as the PDF, as an editable Word document.

It takes the very same story the PDF is built from (build_report.make_story), so the text, numbers, tables and
figures can never drift apart. Headings use Word's real heading styles (so the Navigation pane works), tables are
real Word tables, citations and the contents list are internal links, and the header and footer carry the team
number and an automatic page number.

    python report/build_docx.py

Editing the references: each one is an ordinary paragraph under "References". Citations in the text are plain
"[n]" links to the matching entry, so if you add or reorder references, renumber the "[n]" in the text by hand.
"""
import html
import os
import re
import sys
from html.parser import HTMLParser

import qrcode
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.platypus import Image as RLImage
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, Spacer, Table

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_report as br  # noqa: E402

HERE = br.HERE
OUT_DOCX = os.path.join(br.OUT_DIR, br.BASENAME + ".docx")
FONT = "Arial"
LINK = "0B5CAD"
INK = RGBColor(0x1F, 0x1F, 0x1E)
SOFT = RGBColor(0x55, 0x54, 0x4F)
ACCENT = RGBColor(0xB4, 0x44, 0x1C)

# ---------------------------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------------------------
_bookmark_id = [100]


def set_font(run_or_style_font, size=None, bold=None, italic=None, color=None):
    f = run_or_style_font
    f.name = FONT
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def force_arial(style):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def add_bookmark(paragraph, name):
    _bookmark_id[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bookmark_id[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bookmark_id[0]))
    paragraph._p.append(start)
    paragraph._p.append(end)


def make_run_element(text, bold=False, italic=False, underline=False, size=None, color=None):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), FONT)
    rpr.append(rf)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rpr.append(c)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add_link(paragraph, text, href, bold=False, italic=False, size=None):
    h = OxmlElement("w:hyperlink")
    if href.startswith("#"):
        h.set(qn("w:anchor"), href[1:])
    else:
        rid = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
        h.set(qn("r:id"), rid)
    h.set(qn("w:history"), "1")
    h.append(make_run_element(text, bold=bold, italic=italic, size=size, color=LINK))
    paragraph._p.append(h)


class Markup(HTMLParser):
    """Turns the small reportlab-style markup used in the report (b, i, u, a, br, font) into runs."""

    def __init__(self, paragraph, size=None, base_bold=False, base_italic=False, color=None):
        super().__init__(convert_charrefs=True)
        self.p, self.size, self.color = paragraph, size, color
        self.bold = [base_bold]
        self.italic = [base_italic]
        self.link = None
        self.link_text = []

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "b":
            self.bold.append(True)
        elif tag == "i":
            self.italic.append(True)
        elif tag in ("a", "link"):
            if a.get("name"):
                add_bookmark(self.p, a["name"])
            elif a.get("href"):
                self.link = a["href"]
                self.link_text = []
        elif tag == "br":
            self.p.add_run().add_break(WD_BREAK.LINE)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag in ("a", "link"):
            self.link = None

    def handle_endtag(self, tag):
        if tag == "b" and len(self.bold) > 1:
            self.bold.pop()
        elif tag == "i" and len(self.italic) > 1:
            self.italic.pop()
        elif tag in ("a", "link") and self.link:
            add_link(self.p, "".join(self.link_text), self.link, bold=self.bold[-1], italic=self.italic[-1], size=self.size)
            self.link = None

    def handle_data(self, data):
        if self.link:
            self.link_text.append(data)
            return
        run = self.p.add_run(data)
        run.bold = self.bold[-1] or None
        run.italic = self.italic[-1] or None
        run.font.name = FONT
        if self.size:
            run.font.size = Pt(self.size)
        if self.color is not None:
            run.font.color.rgb = self.color


def fill(paragraph, markup, size=None, bold=False, italic=False, color=None):
    m = Markup(paragraph, size=size, base_bold=bold, base_italic=italic, color=color)
    m.feed(markup)
    m.close()


def add_field(paragraph, instr, size=10, color=None):
    run = paragraph.add_run()
    set_font(run.font, size=size, color=color)
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def cell_borders(cell, bottom=("single", 4, "D9D8D0"), top=None):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("left", None), ("bottom", bottom), ("right", None)):
        el = OxmlElement(f"w:{edge}")
        if spec:
            el.set(qn("w:val"), spec[0])
            el.set(qn("w:sz"), str(spec[1]))
            el.set(qn("w:color"), spec[2])
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcpr.append(b)


def cell_margins(table, top=45, bottom=45, left=70, right=70):
    tblpr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblpr.append(m)


def row_props(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trpr.append(th)


# ---------------------------------------------------------------------------------------------
# document and styles
# ---------------------------------------------------------------------------------------------
def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.1)
    sec.header_distance, sec.footer_distance = Cm(1.1), Cm(0.9)

    st = doc.styles
    normal = st["Normal"]
    set_font(normal.font, size=11, color=INK)
    force_arial(normal)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    def mk(name, base, size, bold=False, italic=False, color=INK, before=0, after=0, keep_next=False):
        try:
            s = st[name]
        except KeyError:
            s = st.add_style(name, 1)
        s.base_style = st[base]
        set_font(s.font, size=size, bold=bold, italic=italic, color=color)
        force_arial(s)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = keep_next
        return s

    mk("Title", "Normal", 20, bold=True, after=8)
    mk("Heading 1", "Normal", 14, bold=True, before=11, after=4, keep_next=True)
    mk("Heading 2", "Normal", 11, bold=True, before=6, after=2, keep_next=True)
    mk("Heading 3", "Normal", 11, bold=True, italic=True, before=4, after=1, keep_next=True)
    mk("Caption", "Normal", 10, italic=True, color=SOFT, before=3, after=7)
    mk("Reference", "Normal", 10, after=3.5)
    mk("Quote Box", "Normal", 10, after=8)
    mk("Small Note", "Normal", 10, color=SOFT)
    mk("Kicker", "Normal", 10, bold=True, color=ACCENT)
    mk("Subtitle", "Normal", 14, color=SOFT, after=4)
    for n in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        # Word's built-in headings carry theme colours and borders: strip them so they match the PDF.
        ppr = st[n].element.get_or_add_pPr()
        for child in list(ppr):
            if child.tag == qn("w:pBdr"):
                ppr.remove(child)
    st["Reference"].paragraph_format.left_indent = Cm(0.8)
    st["Reference"].paragraph_format.first_line_indent = Cm(-0.8)
    qb = st["Quote Box"]
    qb.paragraph_format.left_indent = Cm(0.35)
    qb.paragraph_format.right_indent = Cm(0.2)
    ppr = qb.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F1EC")
    ppr.append(shd)

    # header and footer
    width = sec.page_width - sec.left_margin - sec.right_margin
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("AI Challenge 2026  |  Getting urgent messages to a human, fast\t")
    set_font(r.font, size=10, color=SOFT)
    r = hp.add_run(f"Team {br.TEAM_NUMBER}  |  Page ")
    set_font(r.font, size=10, color=SOFT)
    add_field(hp, "PAGE", size=10, color=SOFT)
    ppr = hp._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "D9D8D0")):
        bt.set(qn(f"w:{k}"), v)
    pb.append(bt)
    ppr.append(pb)

    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(int(width / 2), WD_TAB_ALIGNMENT.CENTER)
    fp.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run(f"Team {br.TEAM_NUMBER}\t")
    set_font(r.font, size=10, color=SOFT)
    add_link(fp, "Back to contents", "#contents", size=10)
    r = fp.add_run("\tPage ")
    set_font(r.font, size=10, color=SOFT)
    add_field(fp, "PAGE", size=10, color=SOFT)
    ppr = fp._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    tp = OxmlElement("w:top")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "D9D8D0")):
        tp.set(qn(f"w:{k}"), v)
    pb.append(tp)
    ppr.append(pb)

    cp = doc.core_properties
    cp.title = "Getting Urgent Messages to a Human, Fast"
    cp.author = f"AI Challenge Team {br.TEAM_NUMBER}"
    cp.subject = "CDU IT Code Fair 2026, AI Challenge, Brief 5"
    return doc


# ---------------------------------------------------------------------------------------------
# flowable -> Word
# ---------------------------------------------------------------------------------------------
STYLE_MAP = {"title": "Title", "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "caption": "Caption", "ref": "Reference",
             "quote": "Quote Box", "small": "Small Note", "kicker": "Kicker", "sub": "Subtitle", "body": "Normal", "bullet": "Normal",
             "cell": "Normal", "cellb": "Normal"}
QR_FILES = []


def qr_png(data):
    path = os.path.join(HERE, "_qr.png")
    qrcode.make(data, box_size=8, border=1).save(path)
    QR_FILES.append(path)
    return path


def add_paragraph(container, flow, in_table=False):
    sname = flow.style.name
    p = container.add_paragraph(style=STYLE_MAP.get(sname, "Normal"))
    text = flow.text
    bullet = getattr(flow, "bulletText", None)
    if bullet:
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.65))
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{bullet}\t")
        set_font(r.font, size=11)
    if in_table:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        fill(p, text, size=10, bold=(sname == "cellb"))
    elif sname in ("h1", "h2", "h3", "title", "caption", "ref", "quote", "small", "kicker", "sub"):
        fill(p, text)
    else:
        fill(p, text, size=11 if not bullet else 11)
    return p


def add_image(container, flow, align=WD_ALIGN_PARAGRAPH.CENTER, width_pt=None):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(flow.filename, width=Pt(width_pt or flow.drawWidth))
    return p


def add_flowable(container, flow, in_table=False):
    if isinstance(flow, br.QR):
        p = container.add_paragraph()
        p.add_run().add_picture(qr_png(flow.data), width=Cm(2.6))
    elif isinstance(flow, Paragraph):
        add_paragraph(container, flow, in_table)
    elif isinstance(flow, RLImage):
        add_image(container, flow)
    elif isinstance(flow, KeepTogether):
        for sub in flow._content:
            add_flowable(container, sub, in_table)
    elif isinstance(flow, Table):
        add_table(container, flow)
    elif isinstance(flow, PageBreak):
        container.add_page_break()
    elif isinstance(flow, Spacer):
        return
    else:
        raise TypeError(f"unhandled flowable {type(flow)}")


def add_table(container, flow):
    rows = flow._cellvalues
    widths = flow._colWidths
    ncols = len(rows[0])
    table = container.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell_margins(table)
    is_header = bool(getattr(flow, "repeatRows", 0))
    for ri, row in enumerate(rows):
        row_props(table.rows[ri], header=is_header and ri == 0)
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.width = Pt(widths[ci])
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            items = val if isinstance(val, (list, tuple)) else [val]
            for it in items:
                if isinstance(it, str):
                    pp = cell.add_paragraph()
                    fill(pp, html.escape(it), size=10)
                else:
                    add_flowable(cell, it, in_table=True)
            if not cell.paragraphs:
                cell.add_paragraph()
            if is_header and ri == 0:
                shade(cell, "F2F1EC")
                cell_borders(cell, bottom=("single", 8, "1F1F1E"))
            else:
                cell_borders(cell)
    # a little air after a table
    container.add_paragraph().paragraph_format.space_after = Pt(0)


def tidy_trailing_blank(doc):
    """Tables are followed by an empty spacer paragraph; make those paragraphs tiny so they do not add pages."""
    for p in doc.paragraphs:
        if not p.text.strip() and not p._p.xpath(".//w:drawing") and not p._p.xpath(".//w:br"):
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(4)
            for r in p.runs:
                r.font.size = Pt(2)


def main():
    br.run_pdf()                                           # settles citation numbers
    story, _ = br.make_story(toc_pages=False)
    doc = new_document()
    body = doc._body
    for flow in story:
        if isinstance(flow, br.Head):
            p = add_paragraph(doc, flow)
            if not p._p.xpath(".//w:bookmarkStart"):
                add_bookmark(p, flow.key)
        else:
            add_flowable(doc, flow)
    tidy_trailing_blank(doc)
    # the first paragraph Word creates for a new document is empty; remove it
    first = doc.paragraphs[0]
    if not first.text.strip() and not first._p.xpath(".//w:drawing"):
        first._p.getparent().remove(first._p)
    doc.save(OUT_DOCX)
    for f in QR_FILES:
        try:
            os.remove(f)
        except OSError:
            pass
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
